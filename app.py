import io
import json
import os
import secrets
import sqlite3
import tempfile
import urllib.error
import urllib.request
from datetime import date, datetime
from functools import wraps
from pathlib import Path

from docx import Document
from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


LEVELS = {"international": ("Международный", 3), "all_russian": ("Всероссийский", 2), "regional": ("Региональный", 1)}
RESULTS = {"first": ("I место / победитель / гран-при / лауреат", 4), "second": ("II место", 3), "third": ("III место", 2), "participant": ("Участник", 1)}
DIRECTIONS = {"research_project": "Исследовательские проекты", "creative": "Творчество", "olympiad": "Олимпиады", "science_technical": "Научно-техническое", "other": "Другое"}
ROLES = {"admin": "Администратор", "editor": "Редактор", "viewer": "Наблюдатель"}
EVENT_TYPES = {"competition": "Конкурс", "olympiad": "Олимпиада", "conference": "Конференция", "festival": "Фестиваль", "grant": "Грант", "other": "Другое"}
FORMATS = {"online": "Онлайн", "offline": "Очно", "hybrid": "Смешанный"}
ALIASES = {"grand prix": "first", "grand_prix": "first", "гран-при": "first", "winner": "first", "победитель": "first", "laureate": "first", "лауреат": "first"}
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "pdf", "doc", "docx", "xls", "xlsx"}


def environment_setting(name, default=""):
    value = os.environ.get(name)
    if value:
        return value
    if os.name == "nt":
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as key:
                return winreg.QueryValueEx(key, name)[0]
        except (FileNotFoundError, OSError):
            pass
    return default


def create_app(test_config=None):
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.environ.get("SECRET_KEY", "change-this-secret-in-production"),
        DATABASE=os.environ.get("DATABASE_PATH", str(Path(app.root_path) / "achievements.db")),
        UPLOAD_FOLDER=os.environ.get("UPLOAD_FOLDER", str(Path(app.root_path) / "uploads")),
        MAX_CONTENT_LENGTH=16 * 1024 * 1024,
        DEEPSEEK_API_KEY=environment_setting("DEEPSEEK_API_KEY"),
        DEEPSEEK_API_URL=environment_setting("DEEPSEEK_API_URL", "https://api.deepseek.com"),
        DEEPSEEK_MODEL=environment_setting("DEEPSEEK_MODEL", "deepseek-flash"),
        PUBLIC_DEMO_MODE=os.environ.get("PUBLIC_DEMO_MODE", "true").lower() in {"1", "true", "yes", "on"},
    )
    if test_config:
        app.config.update(test_config)
    Path(app.config["UPLOAD_FOLDER"]).mkdir(parents=True, exist_ok=True)

    def get_db():
        if "db" not in g:
            g.db = sqlite3.connect(app.config["DATABASE"])
            g.db.row_factory = sqlite3.Row
            g.db.execute("PRAGMA foreign_keys = ON")
        return g.db

    @app.teardown_appcontext
    def close_db(_error=None):
        db = g.pop("db", None)
        if db is not None:
            db.close()

    def audit(action, entity_type="", entity_id=None, details=""):
        db = get_db()
        db.execute("INSERT INTO audit_log(user_id, action, entity_type, entity_id, details, ip) VALUES(?,?,?,?,?,?)",
                   (session.get("user_id"), action, entity_type, entity_id, details, request.remote_addr if request else ""))
        db.commit()

    def init_db():
        db = get_db()
        db.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','editor','viewer')),
                active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS achievements (
                id INTEGER PRIMARY KEY, child_name TEXT NOT NULL, school TEXT NOT NULL,
                class_name TEXT NOT NULL, contest_name TEXT NOT NULL, event_date TEXT NOT NULL,
                level TEXT NOT NULL, result TEXT NOT NULL, direction TEXT NOT NULL,
                mentor_name TEXT NOT NULL, mentor_position TEXT NOT NULL,
                mentor_winner INTEGER NOT NULL DEFAULT 0, mentor_level TEXT, mentor_result TEXT,
                social_notes TEXT, generated_social TEXT, is_demo INTEGER NOT NULL DEFAULT 0,
                created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS attachments (
                id INTEGER PRIMARY KEY, achievement_id INTEGER NOT NULL REFERENCES achievements(id) ON DELETE CASCADE,
                original_name TEXT NOT NULL, stored_name TEXT NOT NULL, uploaded_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id), action TEXT NOT NULL,
                entity_type TEXT, entity_id INTEGER, details TEXT, ip TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS child_accounts (
                id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL,
                child_name TEXT NOT NULL, school TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS opportunities (
                id INTEGER PRIMARY KEY, title TEXT NOT NULL, direction TEXT NOT NULL,
                level TEXT NOT NULL, event_type TEXT NOT NULL, organizer TEXT NOT NULL,
                event_date TEXT NOT NULL, application_deadline TEXT NOT NULL,
                format TEXT NOT NULL, description TEXT NOT NULL, external_url TEXT,
                active INTEGER NOT NULL DEFAULT 1, created_by INTEGER REFERENCES users(id),
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS invitations (
                id INTEGER PRIMARY KEY,
                opportunity_id INTEGER NOT NULL REFERENCES opportunities(id) ON DELETE CASCADE,
                child_account_id INTEGER NOT NULL REFERENCES child_accounts(id) ON DELETE CASCADE,
                status TEXT NOT NULL DEFAULT 'unread' CHECK(status IN ('unread','viewed','interested','declined')),
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, viewed_at TEXT, responded_at TEXT,
                UNIQUE(opportunity_id, child_account_id)
            );
            CREATE INDEX IF NOT EXISTS idx_ach_child ON achievements(child_name);
            CREATE INDEX IF NOT EXISTS idx_ach_mentor ON achievements(mentor_name);
            CREATE INDEX IF NOT EXISTS idx_child_identity ON child_accounts(child_name, school);
            CREATE INDEX IF NOT EXISTS idx_inv_child ON invitations(child_account_id, status);
        """)
        columns = {row[1] for row in db.execute("PRAGMA table_info(achievements)").fetchall()}
        if "generated_social" not in columns:
            db.execute("ALTER TABLE achievements ADD COLUMN generated_social TEXT")
        if not db.execute("SELECT 1 FROM users LIMIT 1").fetchone():
            db.execute("INSERT INTO users(username,password_hash,role) VALUES(?,?,?)",
                       ("admin", generate_password_hash("admin123", method="pbkdf2:sha256"), "admin"))
        if not db.execute("SELECT 1 FROM achievements LIMIT 1").fetchone():
            uid = db.execute("SELECT id FROM users WHERE username='admin'").fetchone()[0]
            demos = [
                ("Иванова Мария Сергеевна", "Школа № 1", "7А", "Юный исследователь", "2026-03-15", "regional", "first", "research_project", "Петрова Анна Викторовна", "учитель биологии", 0, None, None, "Защита проекта об экологии города.", 1, uid),
                ("Смирнов Алексей Игоревич", "Лицей № 2", "9Б", "Всероссийская олимпиада по математике", "2026-02-20", "all_russian", "second", "olympiad", "Соколова Елена Павловна", "учитель математики", 1, "regional", "first", "Поздравляем с высоким результатом!", 1, uid),
                ("Иванова Мария Сергеевна", "Школа № 1", "7А", "Созвездие талантов", "2026-04-04", "international", "third", "creative", "Орлова Ирина Николаевна", "педагог дополнительного образования", 0, None, None, "Яркое выступление в номинации «Вокал».", 1, uid),
            ]
            db.executemany("""INSERT INTO achievements(child_name,school,class_name,contest_name,event_date,level,result,direction,mentor_name,mentor_position,mentor_winner,mentor_level,mentor_result,social_notes,is_demo,created_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", demos)
        if not db.execute("SELECT 1 FROM child_accounts LIMIT 1").fetchone():
            db.execute("INSERT INTO child_accounts(username,password_hash,child_name,school) VALUES(?,?,?,?)",
                       ("student", generate_password_hash("student123", method="pbkdf2:sha256"), "Иванова Мария Сергеевна", "Школа № 1"))
        if not db.execute("SELECT 1 FROM opportunities LIMIT 1").fetchone():
            uid = db.execute("SELECT id FROM users WHERE username='admin'").fetchone()[0]
            db.execute("""INSERT INTO opportunities(title,direction,level,event_type,organizer,event_date,application_deadline,format,description,external_url,active,created_by)
                        VALUES(?,?,?,?,?,?,?,?,?,?,1,?)""",
                       ("Экологические проекты будущего", "research_project", "regional", "competition", "Центр юных исследователей", "2026-10-20", "2026-10-01", "hybrid", "Конкурс исследовательских и проектных работ школьников.", "https://example.org/eco-projects", uid))
        db.execute("""INSERT OR IGNORE INTO invitations(opportunity_id,child_account_id)
                      SELECT o.id,c.id FROM opportunities o JOIN child_accounts c ON c.active=1
                      WHERE o.active=1 AND EXISTS (
                        SELECT 1 FROM achievements a WHERE a.child_name=c.child_name AND a.school=c.school AND a.direction=o.direction
                      )""")
        db.commit()

    with app.app_context():
        init_db()

    def login_required(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not g.user:
                return redirect(url_for("login", next=request.path))
            return view(*args, **kwargs)
        return wrapped

    def public_demo_access():
        return app.config["PUBLIC_DEMO_MODE"] and not g.user

    def demo_or_login_required(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not g.user and not public_demo_access():
                return redirect(url_for("login", next=request.path))
            return view(*args, **kwargs)
        return wrapped

    def roles_required(*roles):
        def deco(view):
            @wraps(view)
            @login_required
            def wrapped(*args, **kwargs):
                if g.user["role"] not in roles:
                    abort(403)
                return view(*args, **kwargs)
            return wrapped
        return deco

    def child_required(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not g.child:
                return redirect(url_for("login", profile="child", next=request.path))
            return view(*args, **kwargs)
        return wrapped

    def invite_matching_children(opportunity_id):
        db = get_db()
        db.execute("""INSERT OR IGNORE INTO invitations(opportunity_id,child_account_id)
                      SELECT o.id,c.id FROM opportunities o JOIN child_accounts c ON c.active=1
                      WHERE o.id=? AND o.active=1 AND EXISTS (
                        SELECT 1 FROM achievements a WHERE a.child_name=c.child_name
                        AND a.school=c.school AND a.direction=o.direction
                      )""", (opportunity_id,))

    def invite_child_to_matching_opportunities(child_id):
        db = get_db()
        db.execute("""INSERT OR IGNORE INTO invitations(opportunity_id,child_account_id)
                      SELECT o.id,c.id FROM opportunities o JOIN child_accounts c ON c.id=? AND c.active=1
                      WHERE o.active=1 AND EXISTS (
                        SELECT 1 FROM achievements a WHERE a.child_name=c.child_name
                        AND a.school=c.school AND a.direction=o.direction
                       )""", (child_id,))

    def invite_identity_to_matching_opportunities(child_name, school):
        db = get_db()
        accounts = db.execute("SELECT id FROM child_accounts WHERE child_name=? AND school=? AND active=1",
                              (child_name, school)).fetchall()
        for account in accounts:
            invite_child_to_matching_opportunities(account["id"])

    @app.before_request
    def load_user_and_csrf():
        g.user = None
        g.child = None
        if session.get("user_id"):
            g.user = get_db().execute("SELECT * FROM users WHERE id=? AND active=1", (session["user_id"],)).fetchone()
            if not g.user:
                session.clear()
        if session.get("child_id"):
            g.child = get_db().execute("SELECT * FROM child_accounts WHERE id=? AND active=1", (session["child_id"],)).fetchone()
            if not g.child:
                session.clear()
        if "csrf_token" not in session:
            session["csrf_token"] = secrets.token_hex(24)
        if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
            token = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token")
            if not secrets.compare_digest(token or "", session["csrf_token"]):
                abort(400, "Недействительный CSRF-токен")

    @app.context_processor
    def template_data():
        return dict(LEVELS=LEVELS, RESULTS=RESULTS, DIRECTIONS=DIRECTIONS, ROLES=ROLES,
                    EVENT_TYPES=EVENT_TYPES, FORMATS=FORMATS,
                    csrf_token=session.get("csrf_token"), score=score, social_text=social_text,
                    public_demo=public_demo_access())

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if g.user:
            return redirect(url_for("dashboard"))
        if g.child:
            return redirect(url_for("cabinet"))
        profile = request.form.get("profile") or request.args.get("profile", "organizer")
        if profile not in {"organizer", "child"}:
            profile = "organizer"
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            next_path = request.args.get("next", "")
            if not next_path.startswith("/") or next_path.startswith("//"):
                next_path = ""
            if profile == "child":
                child = get_db().execute("SELECT * FROM child_accounts WHERE username=? AND active=1", (username,)).fetchone()
                if child and check_password_hash(child["password_hash"], password):
                    session.clear()
                    session["child_id"] = child["id"]
                    session["csrf_token"] = secrets.token_hex(24)
                    audit("child_login", "child_account", child["id"], f"child_account={child['username']}")
                    return redirect(next_path or url_for("cabinet"))
            else:
                user = get_db().execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
                if user and check_password_hash(user["password_hash"], password):
                    session.clear()
                    session["user_id"] = user["id"]
                    session["csrf_token"] = secrets.token_hex(24)
                    audit("login", "user", user["id"], "Успешный вход")
                    return redirect(next_path or url_for("dashboard"))
            flash("Неверный логин, пароль или профиль", "error")
        return render_template("login.html", selected_profile=profile)

    @app.post("/logout")
    @login_required
    def logout():
        audit("logout", "user", g.user["id"])
        session.clear()
        return redirect(url_for("login"))

    @app.route("/cabinet/login", methods=["GET", "POST"])
    def cabinet_login():
        target = url_for("login", profile="child", next=request.args.get("next", ""))
        return redirect(target, code=307 if request.method == "POST" else 302)

    @app.post("/cabinet/logout")
    @child_required
    def cabinet_logout():
        audit("child_logout", "child_account", g.child["id"], f"child_account={g.child['username']}")
        session.clear()
        return redirect(url_for("login", profile="child"))

    @app.get("/cabinet")
    @child_required
    def cabinet():
        db = get_db()
        rows = db.execute("""SELECT * FROM achievements WHERE child_name=? AND school=?
                           ORDER BY event_date DESC,id DESC""", (g.child["child_name"], g.child["school"])).fetchall()
        totals = {key: 0 for key in DIRECTIONS}
        counts = {key: 0 for key in RESULTS}
        for row in rows:
            totals[row["direction"]] += 1
            result_key = ALIASES.get(str(row["result"]).lower(), row["result"])
            if result_key in counts:
                counts[result_key] += 1
        total_score = sum(score(row["level"], row["result"]) for row in rows)
        ratings = db.execute("SELECT child_name,school,level,result FROM achievements").fetchall()
        child_scores = {}
        for row in ratings:
            identity = (row["child_name"], row["school"])
            child_scores[identity] = child_scores.get(identity, 0) + score(row["level"], row["result"])
        ordered_scores = sorted(child_scores.values(), reverse=True)
        ranking = ordered_scores.index(total_score) + 1 if total_score in ordered_scores else None
        invitations = db.execute("""SELECT i.*,o.*,
                                  i.id AS invitation_id,i.status AS invitation_status
                                  FROM invitations i JOIN opportunities o ON o.id=i.opportunity_id
                                  WHERE i.child_account_id=? ORDER BY o.application_deadline,i.id DESC""",
                                 (g.child["id"],)).fetchall()
        unread_count = sum(1 for item in invitations if item["invitation_status"] == "unread")
        if unread_count:
            db.execute("""UPDATE invitations SET status='viewed',viewed_at=CURRENT_TIMESTAMP
                          WHERE child_account_id=? AND status='unread'""", (g.child["id"],))
            db.commit()
        return render_template("cabinet.html", rows=rows, total_score=total_score, counts=counts,
                               profile=totals, ranking=ranking, invitations=invitations,
                               unread_count=unread_count, ranked_children=len(child_scores))

    @app.post("/cabinet/invitations/<int:invitation_id>/respond")
    @child_required
    def invitation_respond(invitation_id):
        status = request.form.get("status")
        if status not in {"interested", "declined"}:
            abort(400)
        db = get_db()
        invitation = db.execute("SELECT * FROM invitations WHERE id=? AND child_account_id=?",
                                (invitation_id, g.child["id"])).fetchone()
        if not invitation:
            abort(404)
        db.execute("""UPDATE invitations SET status=?,viewed_at=COALESCE(viewed_at,CURRENT_TIMESTAMP),
                      responded_at=CURRENT_TIMESTAMP WHERE id=?""", (status, invitation_id))
        db.commit()
        audit("child_invitation_response", "invitation", invitation_id,
              f"child_account={g.child['username']}; status={status}")
        flash("Ответ сохранён", "success")
        return redirect(url_for("cabinet") + "#invitations")

    def filters_sql():
        clauses, params = (["is_demo=1"] if public_demo_access() else []), []
        for field in ("school", "direction", "level"):
            value = request.args.get(field, "").strip()
            if value:
                clauses.append(f"{field}=?")
                params.append(value)
        start, end = request.args.get("date_from", ""), request.args.get("date_to", "")
        if start:
            clauses.append("event_date>=?"); params.append(start)
        if end:
            clauses.append("event_date<=?"); params.append(end)
        return (" WHERE " + " AND ".join(clauses)) if clauses else "", params

    def call_deepseek(system_prompt, user_prompt, temperature=0.2):
        api_key = app.config.get("DEEPSEEK_API_KEY", "").strip()
        if not api_key:
            raise RuntimeError("Ключ DeepSeek API не настроен")
        payload = json.dumps({
            "model": app.config["DEEPSEEK_MODEL"],
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            "temperature": temperature,
            "stream": False,
        }, ensure_ascii=False).encode("utf-8")
        endpoint = app.config["DEEPSEEK_API_URL"].rstrip("/") + "/chat/completions"
        api_request = urllib.request.Request(endpoint, data=payload, method="POST", headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        })
        try:
            with urllib.request.urlopen(api_request, timeout=12) as response:
                result = json.loads(response.read().decode("utf-8"))
            return result["choices"][0]["message"]["content"].strip()
        except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, KeyError, IndexError, json.JSONDecodeError) as exc:
            raise RuntimeError("Сервис генерации временно недоступен") from exc

    def suggest_direction_with_ai(contest_name):
        try:
            answer = call_deepseek(
                "Ты классификатор образовательных мероприятий. Ответь только одним допустимым кодом без пояснений.",
                "Определи направление мероприятия по названию. Допустимые коды: "
                "research_project — научно-исследовательская и проектная деятельность; "
                "creative — творческий конкурс; olympiad — олимпиада; "
                "science_technical — научно-техническое направление; other — прочее.\n"
                f"Название: {contest_name}",
                0,
            ).lower().strip(" `\n\t.,:;")
            if answer in DIRECTIONS:
                return answer, "ai"
        except RuntimeError:
            pass
        return suggest_direction(contest_name), "rules"

    @app.route("/")
    @demo_or_login_required
    def dashboard():
        db = get_db(); where, params = filters_sql()
        rows = db.execute("SELECT * FROM achievements" + where + " ORDER BY event_date DESC", params).fetchall()
        schools = db.execute("SELECT DISTINCT school FROM achievements ORDER BY school").fetchall()
        by_school, by_direction, school_profiles = {}, {}, {}
        for row in rows:
            by_school[row["school"]] = by_school.get(row["school"], 0) + score(row["level"], row["result"])
            by_direction[row["direction"]] = by_direction.get(row["direction"], 0) + 1
            profile = school_profiles.setdefault(row["school"], {key: 0 for key in DIRECTIONS})
            profile[row["direction"]] += 1
        return render_template("dashboard.html", rows=rows, schools=schools, by_school=by_school,
                               by_direction=by_direction, school_profiles=school_profiles,
                               total_score=sum(score(r["level"], r["result"]) for r in rows))

    @app.get("/achievements")
    @demo_or_login_required
    def achievements():
        where, params = filters_sql()
        rows = get_db().execute("SELECT * FROM achievements" + where + " ORDER BY event_date DESC, id DESC", params).fetchall()
        schools = get_db().execute("SELECT DISTINCT school FROM achievements ORDER BY school").fetchall()
        return render_template("achievements.html", rows=rows, schools=schools)

    def achievement_values(existing=None):
        fields = ["child_name", "school", "class_name", "contest_name", "event_date", "level", "result", "direction", "mentor_name", "mentor_position", "mentor_level", "mentor_result", "social_notes"]
        data = {f: request.form.get(f, "").strip() for f in fields}
        data["result"] = ALIASES.get(data["result"].lower(), data["result"])
        if not data["direction"]:
            data["direction"] = suggest_direction_with_ai(data["contest_name"])[0]
        data["mentor_winner"] = 1 if request.form.get("mentor_winner") else 0
        required = ["child_name", "school", "class_name", "contest_name", "event_date", "level", "result", "direction", "mentor_name", "mentor_position"]
        if any(not data[k] for k in required):
            raise ValueError("Заполните все обязательные поля")
        if data["level"] not in LEVELS or data["result"] not in RESULTS or data["direction"] not in DIRECTIONS:
            raise ValueError("Выбрано недопустимое значение")
        try: date.fromisoformat(data["event_date"])
        except ValueError: raise ValueError("Некорректная дата")
        if data["mentor_winner"]:
            if data["mentor_level"] not in LEVELS or data["mentor_result"] not in RESULTS:
                raise ValueError("Укажите уровень и результат конкурса педагога")
        else:
            data["mentor_level"] = data["mentor_result"] = None
        return data

    def checked_uploads():
        checked = []
        for upload in request.files.getlist("attachments"):
            if not upload or not upload.filename:
                continue
            original = Path(upload.filename).name
            ext = original.rsplit(".", 1)[-1].lower() if "." in original else ""
            if ext not in ALLOWED_EXTENSIONS:
                raise ValueError(f"Недопустимый тип файла: {original}")
            clean = secure_filename(original) or f"attachment.{ext}"
            if "." not in clean:
                clean = f"{clean}.{ext}"
            checked.append((upload, clean, ext))
        return checked

    def save_uploads(achievement_id, uploads):
        db = get_db()
        for upload, clean, ext in uploads:
            stored = f"{secrets.token_hex(12)}.{ext}"
            upload.save(Path(app.config["UPLOAD_FOLDER"]) / stored)
            db.execute("INSERT INTO attachments(achievement_id,original_name,stored_name) VALUES(?,?,?)", (achievement_id, clean, stored))
        db.commit()

    @app.route("/achievements/new", methods=["GET", "POST"])
    @roles_required("admin", "editor")
    def achievement_new():
        if request.method == "POST":
            try:
                data = achievement_values()
                uploads = checked_uploads()
                cols = list(data); values = [data[c] for c in cols]
                db = get_db()
                cur = db.execute(f"INSERT INTO achievements({','.join(cols)},created_by) VALUES({','.join('?' for _ in cols)},?)", values + [g.user["id"]])
                invite_identity_to_matching_opportunities(data["child_name"], data["school"])
                db.commit(); save_uploads(cur.lastrowid, uploads)
                audit("create", "achievement", cur.lastrowid, data["contest_name"])
                flash("Достижение добавлено", "success")
                return redirect(url_for("achievement_detail", achievement_id=cur.lastrowid))
            except ValueError as exc:
                flash(str(exc), "error")
        return render_template("achievement_form.html", item=None)

    @app.get("/achievements/<int:achievement_id>")
    @demo_or_login_required
    def achievement_detail(achievement_id):
        query = "SELECT * FROM achievements WHERE id=?"
        if public_demo_access():
            query += " AND is_demo=1"
        item = get_db().execute(query, (achievement_id,)).fetchone()
        if not item: abort(404)
        files = [] if public_demo_access() else get_db().execute(
            "SELECT * FROM attachments WHERE achievement_id=? ORDER BY id", (achievement_id,)
        ).fetchall()
        return render_template("achievement_detail.html", item=item, files=files)

    @app.route("/achievements/<int:achievement_id>/edit", methods=["GET", "POST"])
    @roles_required("admin", "editor")
    def achievement_edit(achievement_id):
        db = get_db(); item = db.execute("SELECT * FROM achievements WHERE id=?", (achievement_id,)).fetchone()
        if not item: abort(404)
        if request.method == "POST":
            try:
                data = achievement_values(item)
                uploads = checked_uploads()
                db.execute("UPDATE achievements SET " + ",".join(f"{k}=?" for k in data) + ",generated_social=NULL,updated_at=CURRENT_TIMESTAMP,is_demo=0 WHERE id=?", list(data.values()) + [achievement_id])
                invite_identity_to_matching_opportunities(data["child_name"], data["school"])
                db.commit(); save_uploads(achievement_id, uploads); audit("update", "achievement", achievement_id, data["contest_name"])
                flash("Изменения сохранены", "success")
                return redirect(url_for("achievement_detail", achievement_id=achievement_id))
            except ValueError as exc: flash(str(exc), "error")
        return render_template("achievement_form.html", item=item)

    @app.post("/achievements/<int:achievement_id>/delete")
    @roles_required("admin", "editor")
    def achievement_delete(achievement_id):
        db = get_db(); item = db.execute("SELECT contest_name FROM achievements WHERE id=?", (achievement_id,)).fetchone()
        if not item: abort(404)
        files = db.execute("SELECT stored_name FROM attachments WHERE achievement_id=?", (achievement_id,)).fetchall()
        db.execute("DELETE FROM achievements WHERE id=?", (achievement_id,)); db.commit()
        for file in files:
            try: (Path(app.config["UPLOAD_FOLDER"]) / file["stored_name"]).unlink(missing_ok=True)
            except OSError: pass
        audit("delete", "achievement", achievement_id, item["contest_name"])
        flash("Запись удалена", "success")
        return redirect(url_for("achievements"))

    @app.get("/attachments/<int:file_id>")
    @login_required
    def attachment(file_id):
        file = get_db().execute("SELECT * FROM attachments WHERE id=?", (file_id,)).fetchone()
        if not file: abort(404)
        return send_file(Path(app.config["UPLOAD_FOLDER"]) / file["stored_name"], as_attachment=True, download_name=file["original_name"])

    @app.get("/api/suggest-direction")
    @login_required
    def direction_api():
        direction, source = suggest_direction_with_ai(request.args.get("contest", ""))
        return jsonify(direction=direction, source=source)

    @app.post("/api/achievements/<int:achievement_id>/generate-social")
    @roles_required("admin", "editor")
    def generate_social(achievement_id):
        db = get_db()
        item = db.execute("SELECT * FROM achievements WHERE id=?", (achievement_id,)).fetchone()
        if not item:
            abort(404)
        prompt = (
            "Составь единый поздравительный текст на русском языке для публикации в социальных сетях. "
            "Текст должен быть доброжелательным, торжественным, грамотным и готовым к публикации. "
            "Упомяни ребёнка, образовательную организацию, класс, конкурс, уровень, результат и педагога-наставника. "
            "Не придумывай факты, награды, цитаты и хэштеги, которых нет в данных. Объём — до 900 знаков.\n\n"
            f"Ребёнок: {item['child_name']}\nОрганизация: {item['school']}\nКласс: {item['class_name']}\n"
            f"Конкурс: {item['contest_name']}\nДата: {item['event_date']}\n"
            f"Уровень: {LEVELS[item['level']][0]}\nРезультат: {RESULTS[item['result']][0]}\n"
            f"Направление: {DIRECTIONS[item['direction']]}\nНаставник: {item['mentor_name']}\n"
            f"Должность наставника: {item['mentor_position']}\nДополнительная информация: {item['social_notes'] or 'нет'}"
        )
        try:
            text = call_deepseek(
                "Ты редактор официальных страниц образовательной организации. Пиши точно по предоставленным данным.",
                prompt,
                0.6,
            )
        except RuntimeError as exc:
            return jsonify(error=str(exc)), 502
        db.execute("UPDATE achievements SET generated_social=?,updated_at=CURRENT_TIMESTAMP WHERE id=?", (text, achievement_id))
        db.commit()
        audit("generate_social", "achievement", achievement_id, app.config["DEEPSEEK_MODEL"])
        return jsonify(text=text)

    @app.get("/children")
    @demo_or_login_required
    def children():
        rows = registry_rows(get_db(), "child", public_demo_access())
        return render_template("registry.html", title="Реестр детей", kind="child", rows=rows)

    @app.get("/children/<path:name>")
    @demo_or_login_required
    def child_detail(name):
        query = "SELECT * FROM achievements WHERE child_name=?"
        if public_demo_access():
            query += " AND is_demo=1"
        rows = get_db().execute(query + " ORDER BY event_date DESC", (name,)).fetchall()
        if not rows: abort(404)
        return render_template("person_detail.html", title=name, rows=rows, kind="child")

    @app.get("/mentors")
    @demo_or_login_required
    def mentors():
        rows = registry_rows(get_db(), "mentor", public_demo_access())
        return render_template("registry.html", title="Реестр педагогов", kind="mentor", rows=rows)

    @app.get("/mentors/<path:name>")
    @demo_or_login_required
    def mentor_detail(name):
        query = "SELECT * FROM achievements WHERE mentor_name=?"
        if public_demo_access():
            query += " AND is_demo=1"
        rows = get_db().execute(query + " ORDER BY event_date DESC", (name,)).fetchall()
        if not rows: abort(404)
        return render_template("person_detail.html", title=name, rows=rows, kind="mentor")

    def excel_export(kind):
        rows = registry_rows(get_db(), kind)
        wb = Workbook(); ws = wb.active; ws.title = "Дети" if kind == "child" else "Педагоги"
        headers = ["ФИО", "Организация / должность", "Достижений", "I место", "II место", "III место", "Участие", "Баллы детей"]
        if kind == "mentor": headers += ["Баллы собственных конкурсов", "Итого"]
        else: headers += ["Итого"]
        ws.append(headers)
        for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="3157A4")
        for row in rows:
            values = [row["name"], row["subtitle"], row["count"], row["first_count"], row["second_count"], row["third_count"], row["participant_count"], row["child_score"]]
            if kind == "mentor": values.append(row["own_score"])
            values.append(row["total_score"]); ws.append(values)
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(45, max(len(str(c.value or "")) for c in col) + 2)
        stream = io.BytesIO(); wb.save(stream); stream.seek(0)
        return send_file(stream, as_attachment=True, download_name=f"{kind}_registry.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.get("/exports/children.xlsx")
    @login_required
    def export_children(): return excel_export("child")

    @app.get("/exports/mentors.xlsx")
    @login_required
    def export_mentors(): return excel_export("mentor")

    @app.get("/achievements/<int:achievement_id>/social.docx")
    @login_required
    def export_social(achievement_id):
        item = get_db().execute("SELECT * FROM achievements WHERE id=?", (achievement_id,)).fetchone()
        if not item: abort(404)
        doc = Document(); doc.add_heading("Поздравляем!", 0)
        for paragraph in social_text(item).split("\n"):
            doc.add_paragraph(paragraph)
        stream = io.BytesIO(); doc.save(stream); stream.seek(0)
        return send_file(stream, as_attachment=True, download_name=f"congratulation_{achievement_id}.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    def opportunity_values():
        fields = ["title", "direction", "level", "event_type", "organizer", "event_date",
                  "application_deadline", "format", "description", "external_url"]
        data = {field: request.form.get(field, "").strip() for field in fields}
        required = [field for field in fields if field != "external_url"]
        if any(not data[field] for field in required):
            raise ValueError("Заполните все обязательные поля")
        if data["direction"] not in DIRECTIONS or data["level"] not in LEVELS:
            raise ValueError("Выбрано недопустимое направление или уровень")
        if data["event_type"] not in EVENT_TYPES or data["format"] not in FORMATS:
            raise ValueError("Выбран недопустимый тип или формат")
        try:
            event_date = date.fromisoformat(data["event_date"])
            deadline = date.fromisoformat(data["application_deadline"])
        except ValueError as exc:
            raise ValueError("Некорректная дата") from exc
        if deadline > event_date:
            raise ValueError("Срок подачи не может быть позже даты мероприятия")
        if data["external_url"] and not data["external_url"].lower().startswith(("http://", "https://")):
            raise ValueError("Внешняя ссылка должна начинаться с http:// или https://")
        data["active"] = 1 if request.form.get("active") else 0
        return data

    @app.get("/opportunities")
    @login_required
    def opportunities():
        rows = get_db().execute("""SELECT o.*,COUNT(i.id) AS invite_count,
                                 SUM(CASE WHEN i.status='interested' THEN 1 ELSE 0 END) AS interested_count
                                 FROM opportunities o LEFT JOIN invitations i ON i.opportunity_id=o.id
                                 GROUP BY o.id ORDER BY o.application_deadline,o.id DESC""").fetchall()
        return render_template("opportunities.html", rows=rows)

    @app.route("/opportunities/new", methods=["GET", "POST"])
    @roles_required("admin", "editor")
    def opportunity_new():
        if request.method == "POST":
            try:
                data = opportunity_values()
                columns = list(data)
                db = get_db()
                cur = db.execute(f"INSERT INTO opportunities({','.join(columns)},created_by) VALUES({','.join('?' for _ in columns)},?)",
                                 [data[column] for column in columns] + [g.user["id"]])
                invite_matching_children(cur.lastrowid)
                db.commit()
                audit("create", "opportunity", cur.lastrowid, data["title"])
                flash("Возможность создана", "success")
                return redirect(url_for("opportunity_detail", opportunity_id=cur.lastrowid))
            except ValueError as exc:
                flash(str(exc), "error")
        return render_template("opportunity_form.html", item=None)

    @app.get("/opportunities/<int:opportunity_id>")
    @login_required
    def opportunity_detail(opportunity_id):
        item = get_db().execute("""SELECT o.*,COUNT(i.id) AS invite_count,
                                 SUM(CASE WHEN i.status='interested' THEN 1 ELSE 0 END) AS interested_count
                                 FROM opportunities o LEFT JOIN invitations i ON i.opportunity_id=o.id
                                 WHERE o.id=? GROUP BY o.id""", (opportunity_id,)).fetchone()
        if not item:
            abort(404)
        return render_template("opportunity_detail.html", item=item)

    @app.route("/opportunities/<int:opportunity_id>/edit", methods=["GET", "POST"])
    @roles_required("admin", "editor")
    def opportunity_edit(opportunity_id):
        db = get_db()
        item = db.execute("SELECT * FROM opportunities WHERE id=?", (opportunity_id,)).fetchone()
        if not item:
            abort(404)
        if request.method == "POST":
            try:
                data = opportunity_values()
                db.execute("UPDATE opportunities SET " + ",".join(f"{key}=?" for key in data) +
                           ",updated_at=CURRENT_TIMESTAMP WHERE id=?", list(data.values()) + [opportunity_id])
                invite_matching_children(opportunity_id)
                db.commit()
                audit("update", "opportunity", opportunity_id, data["title"])
                flash("Возможность обновлена", "success")
                return redirect(url_for("opportunity_detail", opportunity_id=opportunity_id))
            except ValueError as exc:
                flash(str(exc), "error")
        return render_template("opportunity_form.html", item=item)

    @app.route("/admin/child-accounts", methods=["GET", "POST"])
    @roles_required("admin")
    def child_accounts():
        db = get_db()
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            child_name = request.form.get("child_name", "").strip()
            school = request.form.get("school", "").strip()
            if len(username) < 3 or len(password) < 8 or not child_name or not school:
                flash("Укажите логин от 3 символов, пароль от 8 символов, ФИО и организацию", "error")
            elif db.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone():
                flash("Этот логин уже используется сотрудником", "error")
            else:
                try:
                    cur = db.execute("""INSERT INTO child_accounts(username,password_hash,child_name,school)
                                      VALUES(?,?,?,?)""", (username, generate_password_hash(password, method="pbkdf2:sha256"), child_name, school))
                    invite_child_to_matching_opportunities(cur.lastrowid)
                    db.commit()
                    audit("create", "child_account", cur.lastrowid, f"child_account={username}; {child_name}; {school}")
                    flash("Личный кабинет создан", "success")
                except sqlite3.IntegrityError:
                    flash("Такой логин уже существует", "error")
        accounts = db.execute("""SELECT c.*,COUNT(i.id) AS invite_count FROM child_accounts c
                               LEFT JOIN invitations i ON i.child_account_id=c.id
                               GROUP BY c.id ORDER BY c.child_name,c.username""").fetchall()
        identities = db.execute("SELECT DISTINCT child_name,school FROM achievements ORDER BY child_name,school").fetchall()
        return render_template("child_accounts.html", accounts=accounts, identities=identities)

    @app.post("/admin/child-accounts/<int:account_id>/update")
    @roles_required("admin")
    def child_account_update(account_id):
        db = get_db()
        account = db.execute("SELECT * FROM child_accounts WHERE id=?", (account_id,)).fetchone()
        if not account:
            abort(404)
        username = request.form.get("username", "").strip()
        child_name = request.form.get("child_name", "").strip()
        school = request.form.get("school", "").strip()
        password = request.form.get("password", "")
        active = 1 if request.form.get("active") else 0
        if len(username) < 3 or not child_name or not school or (password and len(password) < 8):
            flash("Проверьте логин, ФИО, организацию и новый пароль (не менее 8 символов)", "error")
            return redirect(url_for("child_accounts"))
        if db.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone():
            flash("Этот логин уже используется сотрудником", "error")
            return redirect(url_for("child_accounts"))
        try:
            db.execute("""UPDATE child_accounts SET username=?,child_name=?,school=?,active=?,
                          updated_at=CURRENT_TIMESTAMP WHERE id=?""", (username, child_name, school, active, account_id))
            if password:
                db.execute("UPDATE child_accounts SET password_hash=? WHERE id=?",
                           (generate_password_hash(password, method="pbkdf2:sha256"), account_id))
            invite_child_to_matching_opportunities(account_id)
            db.commit()
            audit("update", "child_account", account_id,
                  f"child_account={username}; active={active}; password_reset={bool(password)}")
            flash("Учётная запись обновлена", "success")
        except sqlite3.IntegrityError:
            db.rollback()
            flash("Такой логин уже существует", "error")
        return redirect(url_for("child_accounts"))

    @app.route("/admin/users", methods=["GET", "POST"])
    @roles_required("admin")
    def users():
        db = get_db()
        if request.method == "POST":
            username = request.form.get("username", "").strip(); password = request.form.get("password", ""); role = request.form.get("role")
            if len(username) < 3 or len(password) < 6 or role not in ROLES:
                flash("Логин от 3 символов, пароль от 6 символов и корректная роль обязательны", "error")
            elif db.execute("SELECT 1 FROM child_accounts WHERE username=?", (username,)).fetchone():
                flash("Этот логин уже используется ребёнком", "error")
            else:
                try:
                    cur = db.execute("INSERT INTO users(username,password_hash,role) VALUES(?,?,?)", (username, generate_password_hash(password, method="pbkdf2:sha256"), role)); db.commit()
                    audit("create", "user", cur.lastrowid, f"{username}: {role}"); flash("Пользователь создан", "success")
                except sqlite3.IntegrityError: flash("Такой логин уже существует", "error")
        return render_template("users.html", users=db.execute("SELECT * FROM users ORDER BY username").fetchall())

    @app.post("/admin/users/<int:user_id>/update")
    @roles_required("admin")
    def user_update(user_id):
        db = get_db(); target = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
        if not target: abort(404)
        role = request.form.get("role"); active = 1 if request.form.get("active") else 0; password = request.form.get("password", "")
        if role not in ROLES or (user_id == g.user["id"] and not active):
            flash("Недопустимое изменение пользователя", "error")
        else:
            db.execute("UPDATE users SET role=?,active=? WHERE id=?", (role, active, user_id))
            if password:
                if len(password) < 6: flash("Пароль должен содержать не менее 6 символов", "error"); return redirect(url_for("users"))
                db.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(password, method="pbkdf2:sha256"), user_id))
            db.commit(); audit("update", "user", user_id, f"role={role}, active={active}"); flash("Пользователь обновлён", "success")
        return redirect(url_for("users"))

    @app.get("/admin/audit")
    @roles_required("admin")
    def audit_log():
        rows = get_db().execute("SELECT a.*,u.username FROM audit_log a LEFT JOIN users u ON u.id=a.user_id ORDER BY a.id DESC LIMIT 500").fetchall()
        return render_template("audit.html", rows=rows)

    @app.get("/admin/backup")
    @roles_required("admin")
    def backup():
        source = get_db(); fd, path = tempfile.mkstemp(suffix=".sqlite3"); os.close(fd)
        target = sqlite3.connect(path); source.backup(target); target.close()
        return send_file(path, as_attachment=True, download_name=f"achievements-backup-{date.today().isoformat()}.sqlite3")

    @app.errorhandler(403)
    def forbidden(_error): return render_template("error.html", code=403, message="Недостаточно прав"), 403
    @app.errorhandler(404)
    def not_found(_error): return render_template("error.html", code=404, message="Страница не найдена"), 404
    @app.errorhandler(413)
    def too_large(_error): return render_template("error.html", code=413, message="Файл слишком большой (максимум 16 МБ)"), 413

    app.get_db = get_db
    return app


def score(level, result):
    return LEVELS.get(level, ("", 0))[1] * RESULTS.get(ALIASES.get(str(result).lower(), result), ("", 0))[1]


def suggest_direction(name):
    text = name.lower()
    groups = [
        ("olympiad", ["олимпиад", "турнир", "математ", "диктант"]),
        ("research_project", ["исслед", "проект", "конференц", "эколог"]),
        ("science_technical", ["технич", "робот", "инженер", "программ", "it", "науч"]),
        ("creative", ["твор", "вокал", "танц", "рисун", "театр", "музык", "искус"]),
    ]
    for direction, words in groups:
        if any(word in text for word in words): return direction
    return "other"


def social_text(row):
    if "generated_social" in row.keys() and row["generated_social"]:
        return row["generated_social"]
    result = RESULTS.get(row["result"], (row["result"], 0))[0]
    level = LEVELS.get(row["level"], (row["level"], 0))[0].lower()
    notes = f"\n{row['social_notes']}" if row["social_notes"] else ""
    return (f"Поздравляем {row['child_name']} ({row['school']}, {row['class_name']})!\n"
            f"{result} в конкурсе «{row['contest_name']}» ({level} уровень), {row['event_date']}.\n"
            f"Наставник: {row['mentor_name']}, {row['mentor_position']}.{notes}\n"
            f"Желаем новых успехов и ярких побед!")


def registry_rows(db, kind, demo_only=False):
    query = "SELECT * FROM achievements"
    if demo_only:
        query += " WHERE is_demo=1"
    rows = db.execute(query + " ORDER BY event_date DESC").fetchall(); grouped = {}
    for row in rows:
        name = row["child_name"] if kind == "child" else row["mentor_name"]
        subtitle = f"{row['school']}, {row['class_name']}" if kind == "child" else row["mentor_position"]
        item = grouped.setdefault(name, {"name": name, "subtitle": subtitle, "count": 0,
                                           "first_count": 0, "second_count": 0, "third_count": 0,
                                           "participant_count": 0, "child_score": 0,
                                           "own_score": 0, "total_score": 0})
        item["count"] += 1; item["child_score"] += score(row["level"], row["result"])
        result_key = ALIASES.get(str(row["result"]).lower(), row["result"])
        if result_key in RESULTS:
            item[f"{result_key}_count"] += 1
        if kind == "mentor" and row["mentor_winner"]:
            item["own_score"] += score(row["mentor_level"], row["mentor_result"])
        item["total_score"] = item["child_score"] + item["own_score"]
    return sorted(grouped.values(), key=lambda x: (-x["total_score"], x["name"]))


app = create_app()

if __name__ == "__main__":
    app.run(host=os.environ.get("HOST", "0.0.0.0"), port=int(os.environ.get("PORT", "5000")), debug=False)
